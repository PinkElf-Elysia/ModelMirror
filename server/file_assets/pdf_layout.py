"""Conservative geometric PDF extraction, called only inside the bounded worker.

This is not a general layout/OCR engine. A page with an unprovable reading
order is retained for diagnostics and explicitly marked degraded.
"""
from __future__ import annotations

import hashlib
import math
import statistics
import unicodedata
from typing import Any

from openpyxl.utils import get_column_letter


class PdfOutputLimit(ValueError):
    pass


def _normalized_line(text: str) -> str:
    return " ".join(unicodedata.normalize("NFKC", text).split())


def _lines(words: list[dict[str, Any]]) -> list[list[dict[str, Any]]]:
    lines: list[list[dict[str, Any]]] = []
    for word in sorted(words, key=lambda w: (w["top"], w["x0"], w["bottom"], w["text"])):
        if lines and abs(word["top"] - lines[-1][0]["top"]) <= 3:
            lines[-1].append(word)
        else:
            lines.append([word])
    return [sorted(line, key=lambda w: (w["x0"], w["x1"], w["text"])) for line in lines]


def _text(words: list[dict[str, Any]]) -> str:
    return "\n".join(" ".join(w["text"] for w in line) for line in _lines(words)).strip()


def _bbox(words: list[dict[str, Any]]) -> list[float]:
    return [min(w["x0"] for w in words), min(w["top"] for w in words),
            max(w["x1"] for w in words), max(w["bottom"] for w in words)]


def _inside(word: dict[str, Any], box: list[float] | tuple[float, ...]) -> bool:
    return (box[0] <= word["x0"] <= word["x1"] <= box[2]
            and box[1] <= word["top"] <= word["bottom"] <= box[3])


def _intersects(a: list[float], b: list[float]) -> bool:
    return min(a[2], b[2]) > max(a[0], b[0]) and min(a[3], b[3]) > max(a[1], b[1])


def _columns(words: list[dict[str, Any]]) -> tuple[list[list[dict[str, Any]]], bool]:
    if not words:
        return [], False
    lines = _lines(words)
    if any(_intersects(_bbox(a), _bbox(b)) for a, b in zip(lines, lines[1:])):
        return [words], True  # Overlapping row bands do not establish reading order.
    gap_size = max(18.0, 3 * statistics.median(w["bottom"] - w["top"] for w in words))
    gaps: list[tuple[float, float]] = []
    for line in lines:
        wide = [(a["x1"], b["x0"]) for a, b in zip(line, line[1:]) if b["x0"] - a["x1"] >= gap_size]
        if len(wide) > 1:
            return [words], True  # Three or more columns / ambiguous sparse layout.
        gaps.extend(wide)
    if not gaps:
        # Overprinted words are not a valid single-column proof either.
        overlap = any(a["x1"] > b["x0"] + 1 for line in lines for a, b in zip(line, line[1:]))
        return [words], overlap
    if len(gaps) < 2:
        return [words], True
    left, right = max(g[0] for g in gaps), min(g[1] for g in gaps)
    if right - left < gap_size:
        return [words], True
    midpoint = (left + right) / 2
    columns = [[], []]
    for line in lines:
        before = [w for w in line if w["x1"] <= midpoint]
        after = [w for w in line if w["x0"] >= midpoint]
        if len(before) + len(after) != len(line):
            return [words], True
        if before and after and after[0]["x0"] - before[-1]["x1"] < gap_size:
            return [words], True
        columns[0].extend(before)
        columns[1].extend(after)
    if not all(columns):
        return [words], True
    # Both columns must occupy the same vertical band, not disjoint indents.
    bounds = [_bbox(column) for column in columns]
    if min(bounds[0][3], bounds[1][3]) <= max(bounds[0][1], bounds[1][1]):
        return [words], True
    return columns, False


def _tables(page: Any, words: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], bool]:
    tables: list[dict[str, Any]] = []
    degraded = False
    for table in page.find_tables(table_settings={"vertical_strategy": "lines_strict", "horizontal_strategy": "lines_strict"}):
        rows = table.extract()
        cells = list(table.cells)
        box = [float(v) for v in table.bbox]
        # Only a complete rectangular grid with >=2 rows and columns is proven.
        valid = bool(rows) and len(rows) >= 2 and len(rows[0]) >= 2 and all(len(r) == len(rows[0]) for r in rows)
        xs = sorted({v for c in cells if c is not None for v in (c[0], c[2])})
        ys = sorted({v for c in cells if c is not None for v in (c[1], c[3])})
        grid = {(x0, y0, x1, y1) for x0, x1 in zip(xs, xs[1:]) for y0, y1 in zip(ys, ys[1:])}
        if not valid or len(xs) != len(rows[0]) + 1 or len(ys) != len(rows) + 1 or set(cells) != grid:
            degraded = True
            continue
        if any(_intersects(box, other["table_bbox"]) for other in tables):
            degraded = True
            continue
        if any(_intersects(box, _bbox([w])) and not _inside(w, box) for w in words):
            degraded = True
            continue
        rendered = []
        for row_index, row in enumerate(rows, 1):
            rendered.append(" | ".join(
                f"{get_column_letter(column_index)}{row_index}: {_normalized_line(str(value or ''))}"
                for column_index, value in enumerate(row, 1)
            ))
        tables.append({"text": "\n".join(rendered), "kind": "table", "table_bbox": box,
                       "row_range": f"A1:{get_column_letter(len(rows[0]))}{len(rows)}"})
    return tables, degraded


def structured_pdf_payload(pdf: Any, *, max_page_characters: int, max_total_characters: int,
                           remove_repeated: bool = True, preserve_tables: bool = True) -> dict[str, Any]:
    pages = []
    extracted = 0
    if not 0 < len(pdf.pages) <= 1000:
        raise PdfOutputLimit("PDF page count exceeds bounded contract")
    for number, page in enumerate(pdf.pages, 1):
        raw_words = page.extract_words(use_text_flow=False, x_tolerance=3, y_tolerance=3)
        # Count before materializing any table rendering or cross-page buffer.
        count = sum(len(str(w.get("text") or "")) + 1 for w in raw_words)
        extracted += count
        if count > max_page_characters or extracted > max_total_characters:
            raise PdfOutputLimit("PDF extracted text exceeds bounded contract")
        words = []
        uncertain = False
        for raw in raw_words:
            word = {key: float(raw[key]) for key in ("x0", "top", "x1", "bottom")}
            if not all(math.isfinite(v) for v in word.values()):
                raise ValueError("Invalid PDF word coordinates")
            word["text"] = str(raw["text"])
            if (not raw.get("upright", True) or not 0 <= word["x0"] < word["x1"] <= float(page.width) + 1
                    or not 0 <= word["top"] < word["bottom"] <= float(page.height) + 1):
                uncertain = True
            words.append(word)
        tables, uncertain_tables = _tables(page, words) if words else ([], False)
        pages.append({"number": number, "height": float(page.height), "words": words,
                      "tables": tables, "uncertain": uncertain or uncertain_tables,
                      "images": list(page.images)})

    operations = []
    if remove_repeated and len(pages) >= 2:
        occurrences: dict[str, list[tuple[int, list[dict[str, Any]]]]] = {}
        for page in pages:
            for line in _lines(page["words"]):
                bounds = _bbox(line)
                if any(_intersects(bounds, table["table_bbox"]) for table in page["tables"]):
                    continue  # Repeated table content is not a removable page edge.
                if bounds[1] > page["height"] * .08 and bounds[3] < page["height"] * .92:
                    continue
                text = _normalized_line(" ".join(w["text"] for w in line))
                if 2 <= len(text) <= 200:
                    occurrences.setdefault(text, []).append((page["number"], line))
        threshold = max(2, math.ceil(len(pages) * .6))
        for text, instances in sorted(occurrences.items()):
            page_ids = sorted({number for number, _ in instances})
            if len(page_ids) < threshold:
                continue
            for number, line in instances:
                ids = {id(w) for w in line}
                pages[number - 1]["words"] = [w for w in pages[number - 1]["words"] if id(w) not in ids]
            operations.append({"code": "remove_repeated_page_edge", "pages": page_ids,
                               "count": len(instances), "line_hash": hashlib.sha256(text.encode()).hexdigest()})

    sections, empty_pages, degraded_pages = [], [], []
    for page in pages:
        number, words, tables = page["number"], page["words"], page["tables"]
        if not words:
            empty_pages.append(number)
            continue
        # Tables are only structured when their complete geometry is proved.
        outside = [w for w in words if not any(_inside(w, table["table_bbox"]) for table in tables)]
        columns, ambiguous = _columns(outside)
        degraded = page["uncertain"] or ambiguous or bool(page["images"])
        if tables and len(columns) > 1:
            degraded = True  # Mixed columns + spanning table order is not proved.
        if tables and not preserve_tables:
            degraded = True
        if degraded:
            degraded_pages.append(number)
            sections.append({"text": _text(words), "page": number, "kind": "page", "layout_status": "degraded"})
            continue
        pieces = [{"text": _text(column), "page": number, "kind": "page",
                   "column": index if len(columns) > 1 else None, "bbox": _bbox(column)}
                  for index, column in enumerate(columns, 1) if column]
        pieces.extend({**table, "page": number, "bbox": table["table_bbox"]} for table in tables)
        if tables:
            # Split text around tables by vertical position so a title and
            # trailing paragraph cannot straddle the table in one block.
            pieces = [{"text": _text(line), "page": number, "kind": "page", "bbox": _bbox(line)}
                      for line in _lines(outside)] + [{**t, "page": number, "bbox": t["table_bbox"]} for t in tables]
            pieces.sort(key=lambda p: (p["bbox"][1], p["bbox"][0]))
            if any(a["bbox"][3] > b["bbox"][1] for a, b in zip(pieces, pieces[1:])):
                degraded_pages.append(number)
                pieces = [{"text": _text(words), "page": number, "kind": "page", "layout_status": "degraded"}]
        sections.extend(pieces)
    retained = sum(len(s["text"]) for s in sections)
    page_sizes: dict[int, int] = {}
    for section in sections:
        page_sizes[section["page"]] = page_sizes.get(section["page"], 0) + len(section["text"])
    if retained > max_total_characters or any(size > max_page_characters for size in page_sizes.values()):
        raise PdfOutputLimit("PDF rendered structure exceeds bounded contract")
    return {"format": "pdf", "sections": sections, "warnings": ["rag_layout_degraded"] if degraded_pages else [],
            "extracted_chars": extracted, "truncated": False, "page_count": len(pages),
            "empty_pages": empty_pages, "degraded_pages": degraded_pages,
            "layout_status": "degraded" if degraded_pages else "verified", "operations": operations}
